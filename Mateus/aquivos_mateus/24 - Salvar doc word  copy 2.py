import pandas as pd
import os
from docx import Document
from comtypes.client import CreateObject

def convert_to_pdf(doc_path, pdf_path):
    word = CreateObject('Word.Application')
    doc = word.Documents.Open(doc_path)
    doc.SaveAs(pdf_path, FileFormat=17)
    doc.Close()
    word.Quit()

def replace_placeholders(doc, data):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            original_text = run.text
            for key, value in data.items():
                placeholder = f'{{{key}}}'
                if placeholder in original_text:
                    run.text = original_text.replace(placeholder, str(value))

def main():
    csv_path = r'C:\Users\MSI Pulse\Documents\Chatpdf\Effecti\Boletim Resumo - Info Excel\Resumo_pronto.csv'
    df = pd.read_csv(csv_path)
    output_directory = r'C:\Users\MSI Pulse\Documents\Chatpdf\Effecti\Boletim Resumo - Envio'
    modelo_path = r'C:\Users\MSI Pulse\Documents\Chatpdf\Effecti\Modelo Excel Oportunidades\MODELO - RESUMO PRONTO.docx'

    for index, row in df.iterrows():
        doc = Document(modelo_path)
        data = {'ID': str(row['ID'])}
        for key in row.index:
            if key.startswith('UF_CRM_'):
                data[key] = str(row[key])
        
        replace_placeholders(doc, data)

        company_name = row['COMPANY_NAME'].strip()
        sanitized_company_name = ''.join(c for c in company_name if c.isalnum() or c in ' _-')
        company_dir = os.path.join(output_directory, sanitized_company_name)
        os.makedirs(company_dir, exist_ok=True)

        # Os arquivos serão nomeados apenas com a ID
        word_file_path = os.path.join(company_dir, f"{row['ID']}.docx")
        pdf_file_path = os.path.join(company_dir, f"{row['ID']}.pdf")
        doc.save(word_file_path)

        # Convertendo para PDF
        convert_to_pdf(word_file_path, pdf_file_path)

    print("Documentos criados e convertidos para PDF com sucesso.")

if __name__ == "__main__":
    main()
